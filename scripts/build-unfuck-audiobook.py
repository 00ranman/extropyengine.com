#!/usr/bin/env python3
"""Audiobook manuscript for ElevenReader.

DOCX: Heading 1 on every chapter so chapters land in the publishing editor.
TXT: one paragraph per line so wrap is not a pause.

Speech rules:
- sentence case on titles (title case makes Eleven hitch before caps)
- em dashes become commas (a dash is a pause; only keep a pause when we want one)
- meters and glyphs spoken, not spelled like code
- no table of contents (ElevenReader already navigates by chapter)
"""
from __future__ import annotations

import html
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path("/tmp/extropyengine.com")
SRC = ROOT / "public/docs/unfuck-the-world.html"
OUT_TXT = ROOT / "public/docs/Unfuck_The_World_For_A_Dollar.elevenreader.txt"
OUT_DOCX = ROOT / "public/docs/Unfuck_The_World_For_A_Dollar.elevenreader.docx"
ART = Path("/workspace/artifacts")

TITLES = [
    "A note on this edition",
    "Preface: The Revolution They Botched",
    "PART I: THE DIAGNOSIS",
    "PART II: THE THESIS",
    "PART III: THE ENGINE",
    "PART IV: THE DEFENSE",
    "PART V: THE ECOSYSTEM",
    "PART VI: THE CALL",
    "Chapter 1: The System Is Rigged — A Systems-Level Autopsy",
    "Chapter 2: The Big Lie of Scarcity — How They Weaponized Nothing",
    "Chapter 3: Comfort as Control — The Game Theory of Compliance",
    "Chapter 4: Value Is Not Money — The Measurement Revolution",
    "Chapter 5: Entropy Reduction as the Unit of All Value",
    "Chapter 6: The Contribution Economy — Tracking What Actually Matters",
    "Chapter 7: Abundance by Design — Why Scarcity Was Always a Choice",
    "Chapter 8: The XP Formula — Where Physics Meets Economics",
    "Chapter 9: The Eight Domains of Entropy — A Taxonomy of Real Value",
    "Chapter 10: The Core Loop — How Claims Become Truth",
    "Chapter 11: The DAG Substrate — Why Not Blockchain",
    "Chapter 12: DFAOs — Fractal Organizations from MicroTeams to Civilization",
    "Chapter 13: Records and Meters — Nothing You Can Bag",
    "Chapter 14: Temporal Mechanics — Time as a System Variable",
    "Chapter 15: Antifragility — Why Every Attack Makes It Stronger",
    "Chapter 16: The Human Psychology Machine — Weaponizing What Makes Us Weird",
    "Chapter 17: Game Theory and Nash Equilibria — The Math of Inevitability",
    "Chapter 18: Sybil Resistance and Adversarial Modeling — Breaking the Breakers",
    "Chapter 19: The Parallel Economy in Practice — From Households to Nations",
    "Chapter 20: The Extropy Ecosystem — Software, Hardware, Music, and Mind",
    "Chapter 21: Emergent Sovereignty — Self-Validating Systems for a Post-Trust World",
    "Chapter 22: Universal Time — Hydrogen-Anchored Temporal Infrastructure",
    "Chapter 23: One Dollar — The Proof of Concept That Changes Everything",
    "Chapter 24: The Inevitable — Why This Cannot Be Stopped",
    "Appendix A: Formula Reference",
    "Appendix B: Routing, Not Purchasing Power",
    "Appendix C: Common Misreadings",
    "Appendix D: Records and Meters",
]

SKIP_HEADS = {
    "UNF*CK THE WORLD",
    "UNF*CK THE WORLD FOR A DOLLAR",
    "Contents",
    "TABLE OF CONTENTS",
}

# Recap slugs from the engine overview, not real chapters.
RECAP = re.compile(
    r"^(Chapter (8|9|10|11|12|13|14): ("
    r"The XP Formula — XP|"
    r"The Eight Domains —|"
    r"The Core Loop — OPEN|"
    r"The DAG Substrate — A|"
    r"DFAOs — Fractal organizations|"
    r"Records and Meters — XP,|"
    r"Temporal Mechanics — Time as a first-class"
    r"))"
)

KEEP_CAPS = {
    "I",
    "I'm",
    "I'll",
    "I've",
    "I'd",
    "Randall",
    "Gossit",
    "Gossett",
    "Extropy",
    "Eckstrophy",
    "Bitcoin",
    "Ethereum",
    "Nash",
    "Sybil",
    "English",
    "America",
    "American",
    "Mars",
    "God",
    "Okay",
}


def peel(block: str) -> tuple[str, str]:
    s = re.sub(r"\s+", " ", block).strip()
    if not s:
        return "", ""
    if s.startswith("PART III: THE ENGINE"):
        rest = s[len("PART III: THE ENGINE") :].strip()
        rest = re.sub(r"^A note before we dive in:\s*", "A note before we dive in. ", rest)
        return "PART III: THE ENGINE", rest
    if s.startswith("PART IV: THE DEFENSE"):
        rest = s[len("PART IV: THE DEFENSE") :].strip()
        rest = re.sub(r"^Or:\s*", "", rest)
        return "PART IV: THE DEFENSE", rest
    if s.startswith("PART V: THE ECOSYSTEM"):
        rest = s[len("PART V: THE ECOSYSTEM") :].strip()
        rest = re.sub(r"^Or:\s*", "", rest)
        return "PART V: THE ECOSYSTEM", rest
    if s.startswith("PART VI: THE CALL"):
        rest = s[len("PART VI: THE CALL") :].strip()
        rest = re.sub(r"^Or:\s*", "", rest)
        return "PART VI: THE CALL", rest
    for title in sorted(TITLES, key=len, reverse=True):
        if s.startswith(title):
            return title, s[len(title) :].strip()
    if s.startswith("Chapter 1: Contributors"):
        return "", s
    if RECAP.match(s):
        return "", s
    return "", s


def sentence_case_title(title: str) -> str:
    t = title.replace("PART I:", "Part 1.").replace("PART II:", "Part 2.")
    t = t.replace("PART III:", "Part 3.").replace("PART IV:", "Part 4.")
    t = t.replace("PART V:", "Part 5.").replace("PART VI:", "Part 6.")
    t = t.replace("Appendix A:", "Appendix A.").replace("Appendix B:", "Appendix B.")
    t = t.replace("Appendix C:", "Appendix C.").replace("Appendix D:", "Appendix D.")
    t = re.sub(r"^(Chapter \d+):", r"\1.", t)
    t = re.sub(r"^Preface:", "Preface.", t)
    t = t.replace("—", ",").replace("–", ",")
    t = re.sub(r"\s+", " ", t).strip()
    words = t.split(" ")
    out = []
    for i, w in enumerate(words):
        if not w:
            continue
        core = re.sub(r"[^A-Za-z']", "", w)
        if i == 0:
            out.append(w[0].upper() + w[1:] if w[:1].isalpha() else w)
            continue
        if core in KEEP_CAPS:
            out.append(w)
            continue
        if i and re.sub(r"[^A-Za-z']", "", words[i - 1]) == "Appendix" and core in {"A", "B", "C", "D"}:
            out.append(w)
            continue
        out.append(w.lower())
    s = " ".join(out)
    s = re.sub(r"([.!?]\s+)([a-z])", lambda m: m.group(1) + m.group(2).upper(), s)
    s = s.replace("x p formula", "X P formula")
    s = s.replace("dfa os", "dee-fows").replace("dfaos", "dee-fows")
    return s


def speak(text: str) -> str:
    t = text
    t = t.replace("Unf*ck", "unfuck").replace("UNF*CK", "unfuck").replace("Unf*cking", "unfucking")
    t = t.replace("unf*ck", "unfuck")
    t = t.replace("Extropy", "Eckstrophy").replace("extropy", "eckstrophy")
    t = t.replace("Lladnaros", "Yadnaros")
    t = t.replace("Gossett", "Gossit")
    t = t.replace("MicroTeams", "micro teams").replace("Micro-Teams", "micro teams")
    t = t.replace("TypeScript", "Type Script")
    t = t.replace("JavaScript", "Java Script")
    t = t.replace("GitHub", "Git Hub")
    t = t.replace("Web3", "web three").replace("web3", "web three")
    t = t.replace("AIaugmented", "A I augmented")
    t = t.replace("w · E", "w dot E")
    t = t.replace("(w · E)", "(w dot E)")
    t = re.sub(r"\bx\*", "x star", t)
    t = t.replace("→", " then ")
    t = t.replace("×", " times ")
    t = t.replace("·", ", ")
    t = t.replace("ΔS", " delta S ")
    t = t.replace("Tₛ", " T sub s ")
    t = t.replace("H_cap", " H cap ")
    t = t.replace("ℱ", " script F ")
    t = t.replace("κ", " kappa ")
    t = t.replace("β", " beta ")
    t = t.replace("λ", " lambda ")
    t = t.replace("—", ", ")
    t = t.replace("–", ", ")
    t = t.replace("…", ".")
    t = t.replace("...", ".")
    t = t.replace("•", ", ")
    t = re.sub(r"\bDFAOs\b", "dee-fows", t)
    t = re.sub(r"\bDFAO\b", "dee-fow", t)
    t = re.sub(r"\bDAOs\b", "dows", t)
    t = re.sub(r"\bDAO\b", "dow", t)
    t = re.sub(r"\bDAGs\b", "dags", t)
    t = re.sub(r"\bDAG\b", "dag", t)
    t = re.sub(r"\bZKP\b", "zero knowledge proof", t)
    t = re.sub(r"\bDIDs\b", "D I Ds", t)
    t = re.sub(r"\bDID\b", "D I D", t)
    t = re.sub(r"\bNFTs\b", "N F Ts", t)
    t = re.sub(r"\bNFT\b", "N F T", t)
    t = re.sub(r"\bGDP\b", "G D P", t)
    t = re.sub(r"\bAPI\b", "A P I", t)
    t = re.sub(r"\bHOA\b", "H O A", t)
    t = re.sub(r"\bPTO\b", "P T O", t)
    t = re.sub(r"\bXP\b", "X P", t)
    t = re.sub(r"\bCT\b", "C T", t)
    t = re.sub(r"\bEP\b", "E P", t)
    t = re.sub(r"\bDT\b", "D T", t)
    t = re.sub(r"\bIT\b", "I T", t)
    t = re.sub(r"\bCAT\b", "cat", t)
    t = re.sub(r"\blog\(1/", "log of one over ", t)
    t = re.sub(r"log of one over\s+T sub s\s*\)", "log of one over T sub s.", t)
    t = t.replace("•", ", ")
    t = re.sub(r"\bX P\s*=\s*", "X P equals ", t)
    t = re.sub(r"\s+=\s+", " equals ", t)
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\s+,", ",", t)
    t = re.sub(r",{2,}", ",", t)
    t = re.sub(r"\s+\.", ".", t)
    t = t.replace(" ,", ",")
    return t.strip()


def html_blocks(raw: str) -> list[tuple[str, str]]:
    raw = re.sub(r"<style[\s\S]*?</style>", "", raw)
    raw = re.sub(r"<script[\s\S]*?</script>", "", raw)
    raw = re.sub(r"<h1[^>]*>", "\n\n@@@H1@@@", raw)
    raw = re.sub(r"<h2[^>]*>", "\n\n@@@H2@@@", raw)
    raw = re.sub(r"</h[12]>", "\n\n", raw)
    raw = re.sub(r"<p[^>]*>", "\n\n", raw)
    raw = re.sub(r"<br\s*/?>", " ", raw)
    raw = re.sub(r"<li[^>]*>", " ", raw)
    raw = re.sub(r"<[^>]+>", " ", raw)
    raw = html.unescape(raw)
    raw = raw.replace("\xa0", " ")
    chunks = re.split(r"\n\s*\n", raw)
    out: list[tuple[str, str]] = []
    skipping_toc = False
    seen: set[str] = set()
    for chunk in chunks:
        s = re.sub(r"\s+", " ", chunk).strip()
        if not s:
            continue
        kind = "p"
        if s.startswith("@@@H1@@@"):
            kind, s = "h1", s[len("@@@H1@@@") :].strip()
        elif s.startswith("@@@H2@@@"):
            kind, s = "h2", s[len("@@@H2@@@") :].strip()
        if skipping_toc:
            if s.startswith("PART I:") or s.startswith("Preface:"):
                skipping_toc = False
            else:
                continue
        if kind in {"h1", "h2"}:
            title, rest = peel(s)
            head = title or s
            if head in SKIP_HEADS or head.startswith("UNF"):
                if head.lower().startswith("content"):
                    skipping_toc = True
                continue
            if head.lower().startswith("content"):
                skipping_toc = True
                continue
            if RECAP.match(head) or (not title and s.startswith("Chapter 1: Contributors")):
                if rest:
                    out.append(("p", rest))
                continue
            if title in seen and title.startswith("Chapter "):
                if rest:
                    out.append(("p", rest))
                continue
            if title:
                seen.add(title)
                out.append(("h", title))
                if rest:
                    out.append(("p", rest))
            else:
                # untitled heading leftover
                out.append(("p", s))
            continue
        out.append(("p", s))
    while out and out[0][0] != "h":
        out.pop(0)
    return out


def tidy_paras(blocks: list[tuple[str, str]]) -> list[tuple[str, str]]:
    cleaned: list[tuple[str, str]] = []
    for kind, text in blocks:
        if kind == "h":
            cleaned.append(("h", sentence_case_title(speak(text))))
            continue
        t = speak(text)
        if not t or t.lower() in {"the book", "contents"}:
            continue
        if re.fullmatch(r"[\d\s\.,]+", t) and len(t) < 40:
            continue
        if cleaned and cleaned[-1][0] == "p" and t[:1].islower():
            cleaned[-1] = ("p", cleaned[-1][1] + " " + t)
            continue
        cleaned.append(("p", t))
    return cleaned


def write_txt(blocks: list[tuple[str, str]], path: Path) -> None:
    lines = [
        "Unfuck the world for a dollar.",
        "",
        "By Randall Gossit.",
        "",
    ]
    for kind, text in blocks:
        if kind == "h":
            lines.append("")
            lines.append(text)
            lines.append("")
        else:
            lines.append(text)
            lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def set_run_font(run, name: str, size: Pt, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = size
    run.bold = bold
    run.font.color.rgb = RGBColor(0x1A, 0x16, 0x12)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = r.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:cs"), name)


def write_docx(blocks: list[tuple[str, str]], path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True

    for style_name, size, bold in (("Normal", 12, False), ("Heading 1", 16, True)):
        st = doc.styles[style_name]
        st.font.name = "Times New Roman"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0x1A, 0x16, 0x12)
        st.font.all_caps = False
        st.font.small_caps = False

    def add_heading(text: str) -> None:
        p = doc.add_paragraph()
        p.style = doc.styles["Heading 1"]
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        set_run_font(run, "Times New Roman", Pt(16), True)

    def add_body(text: str) -> None:
        p = doc.add_paragraph()
        p.style = doc.styles["Normal"]
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(text)
        set_run_font(run, "Times New Roman", Pt(12), False)

    add_heading("Unfuck the world for a dollar.")
    add_body("By Randall Gossit.")
    for kind, text in blocks:
        if kind == "h":
            add_heading(text)
        else:
            add_body(text)
    doc.save(path)


def main() -> None:
    raw = SRC.read_text(encoding="utf-8")
    blocks = tidy_paras(html_blocks(raw))
    heads = [t for k, t in blocks if k == "h"]
    OUT_TXT.parent.mkdir(parents=True, exist_ok=True)
    ART.mkdir(parents=True, exist_ok=True)
    write_txt(blocks, OUT_TXT)
    write_docx(blocks, OUT_DOCX)
    for dest in (ART / OUT_TXT.name, ART / OUT_DOCX.name):
        dest.write_bytes((OUT_TXT if dest.suffix == ".txt" else OUT_DOCX).read_bytes())
    print("txt", OUT_TXT.stat().st_size, "docx", OUT_DOCX.stat().st_size)
    print("headings", len(heads))
    for h in heads:
        print(" -", h)


if __name__ == "__main__":
    main()
